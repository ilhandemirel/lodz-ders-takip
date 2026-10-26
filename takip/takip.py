#!/usr/bin/env python3
"""İlhan'ın Łódź ders programı takipçisi.

Kullanım:
  python3 takip.py --events events.json --config config.json --template template.html \
                   [--prev snapshot.json] --out OUTDIR [--today YYYY-MM-DD]

Çıktılar (OUTDIR):
  lodz-ders-programi.ics, lodz-ders-programi.xlsx, page.html, snapshot.json,
  rapor.txt (Türkçe değişiklik raporu), sonuc.json (makine için özet)
"""
import argparse, json, itertools, datetime as dt, re, os, sys
from collections import defaultdict

ap = argparse.ArgumentParser()
ap.add_argument('--events', required=True)
ap.add_argument('--config', required=True)
ap.add_argument('--template', required=True)
ap.add_argument('--prev')
ap.add_argument('--out', required=True)
ap.add_argument('--today')
a = ap.parse_args()

raw = json.load(open(a.events, encoding='utf-8'))
cfg = json.load(open(a.config, encoding='utf-8'))
prev = json.load(open(a.prev, encoding='utf-8')) if a.prev and os.path.exists(a.prev) else None
os.makedirs(a.out, exist_ok=True)
today = dt.date.fromisoformat(a.today) if a.today else dt.date.today()
BASE = dt.date.fromisoformat(cfg['base_monday'])
START = dt.date.fromisoformat(cfg['semester_start'])
C = cfg['courses']
GUN = ['Pazartesi', 'Salı', 'Çarşamba', 'Perşembe', 'Cuma', 'Cumartesi', 'Pazar']
AY = ['', 'Oca', 'Şub', 'Mar', 'Nis', 'May', 'Haz', 'Tem', 'Ağu', 'Eyl', 'Eki', 'Kas', 'Ara']
hm = lambda m: f"{m // 60:02d}:{m % 60:02d}"
GK = ['Pzt', 'Sal', 'Çar', 'Per', 'Cum', 'Cmt', 'Paz']
tr_date = lambda d: f"{d.day} {AY[d.month]} {GK[d.weekday()]}"
NOISE = re.compile(r'(start from|weeks?\b|odd|even|make-up)', re.I)
clean = lambda n: '' if NOISE.search(n or '') else (n or '').strip()

def room(r, n):
    r = (r or '').strip()
    if not r:
        return 'Yer: hocadan öğren (salon yok)'
    parts = []
    for x in r.split(' / '):
        m = re.match(r'BUILDING (.+?) ROOM (.+)', x)
        if m:
            parts.append(f"{m.group(1)} · {m.group(2)}")
        elif re.match(r'IFE-(\w+)', x):
            parts.append('IFE (A-16) · ' + x.split('-', 1)[1])
        else:
            parts.append(x)
    return ' / '.join(parts)

skip = {(s['id'], s['date']): s['reason'] for s in cfg.get('skip', [])}

# --- oturumlar ---
sessions = []
events = []
for e in raw['events']:
    if e.get('x'):
        continue
    w = list(e['w'])
    for i, f in enumerate(w):
        if f != '1':
            continue
        day = BASE + dt.timedelta(days=7 * i + e['d'] - 1)
        if day < START:
            continue
        if (e['id'], day.isoformat()) in skip:
            w[i] = '0'
            continue
        sessions.append(dict(id=e['id'], date=day.isoformat(), s=e['s'], e=e['s'] + e['u'], k=e['k'], c=C[e['k']],
                             t=e['t'], g=e['g'], cel=e['cel'], r=room(e['r'], e['n']), p=e['p'], n=clean(e['n'])))
    events.append(dict(e, w=''.join(w)))
sessions.sort(key=lambda x: (x['date'], x['s']))

# --- değişiklikler ---
def key(x): return (x['id'], x['date'])
def label(x):
    d = dt.date.fromisoformat(x['date'])
    grp = '' if x['g'] == 'herkes' else f" ({x['g']})"
    note = f" – {x['n']}" if x['n'] and x['t'] != 'Lecture' else ''
    return f"{tr_date(d)} {hm(x['s'])}–{hm(x['e'])} {x['c']} {x['t']}{grp}{note}, {x['r']}"

changes = []   # (typ, date, start, text)
raw_ch = []    # (typ, id, desc, session)
if prev is not None:
    P = {key(x): x for x in prev['sessions'] if dt.date.fromisoformat(x['date']) >= today}
    N = {key(x): x for x in sessions if dt.date.fromisoformat(x['date']) >= today}
    removed = [P[k] for k in P if k not in N]
    added = [N[k] for k in N if k not in P]
    for k in P.keys() & N.keys():
        o, n = P[k], N[k]
        diffs = []
        if (o['s'], o['e']) != (n['s'], n['e']):
            diffs.append(f"saat {hm(o['s'])}–{hm(o['e'])} → {hm(n['s'])}–{hm(n['e'])}")
        if o['r'] != n['r']:
            diffs.append(f"salon {o['r']} → {n['r']}")
        if o['g'] != n['g']:
            diffs.append(f"grup {o['g']} → {n['g']}")
        if diffs:
            raw_ch.append(('DEĞİŞTİ', n['id'], '; '.join(diffs), n))
    used = set()
    for r_ in removed:
        mate = next((ad for ad in added if key(ad) not in used and ad['date'] == r_['date'] and ad['k'] == r_['k'] and ad['t'] == r_['t']), None)
        if mate:
            used.add(key(mate))
            raw_ch.append(('DEĞİŞTİ', mate['id'], f"önceden {hm(r_['s'])}–{hm(r_['e'])}, {r_['r']}", mate))
        else:
            raw_ch.append(('İPTAL', r_['id'], '', r_))
    for ad in added:
        if key(ad) not in used:
            raw_ch.append(('YENİ', ad['id'], '', ad))
    grp = defaultdict(list)
    for typ, eid, desc, x in raw_ch:
        grp[(typ, eid, desc, x['s'], x['e'], x['r'])].append(x)
    for (typ, eid, desc, s_, e_, r_), xs in grp.items():
        xs.sort(key=lambda z: z['date']); x = xs[0]
        dates = [dt.date.fromisoformat(z['date']) for z in xs]
        g = '' if x['g'] == 'herkes' else f" ({x['g']})"
        note = f" – {x['n']}" if x['n'] else ''
        if len(xs) == 1:
            when = f"{tr_date(dates[0])} {hm(s_)}–{hm(e_)}"
        else:
            when = f"{GK[dates[0].weekday()]} {hm(s_)}–{hm(e_)}, {len(xs)} tarih: " + ', '.join(f"{d.day} {AY[d.month]}" for d in dates)
        txt = f"{x['c']} {x['t']}{g}{note} | {when} | {r_}" + (f"  [{desc}]" if desc else '')
        changes.append((typ, xs[0]['date'], s_, txt))
    changes.sort(key=lambda c: (c[1], c[2]))

# --- çakışmalar (bugünden sonrası) ---
fut = [x for x in sessions if dt.date.fromisoformat(x['date']) >= today]
conflicts = [(x, y) for x, y in itertools.combinations(fut, 2) if x['date'] == y['date'] and x['s'] < y['e'] and y['s'] < x['e']]
tight = [(x, y) for x in fut for y in fut if x is not y and x['date'] == y['date'] and 0 <= y['s'] - x['e'] < 15 and x['r'].split(' ·')[0] != y['r'].split(' ·')[0]]

# --- rapor ---
L = []
if prev is None:
    L.append('İlk çalıştırma: değişiklik karşılaştırması yapılmadı, mevcut program kaydedildi.')
elif changes:
    L.append(f"Celcat'ta senin derslerinde {len(changes)} değişiklik var:")
    for typ, _, _, txt in changes:
        L.append(f"- {typ}: {txt}")
else:
    L.append('Senin derslerinde değişiklik yok.')
if conflicts:
    L.append('')
    L.append('ÇAKIŞMALAR:')
    for x, y in conflicts:
        L.append(f"- {label(x)}  ⟷  {hm(y['s'])}–{hm(y['e'])} {y['c']} {y['t']}")
if tight and (changes or prev is None):
    L.append('')
    L.append('Arası 15 dakikadan az, farklı binada olan dersler:')
    for x, y in tight:
        L.append(f"- {tr_date(dt.date.fromisoformat(x['date']))}: {hm(x['e'])} {x['c']} biter → {hm(y['s'])} {y['c']} ({y['r']})")
report = '\n'.join(L)
open(os.path.join(a.out, 'rapor.txt'), 'w', encoding='utf-8').write(report + '\n')

# --- snapshot ---
json.dump(dict(fetched=raw.get('fetched'), lastChangeAll=raw.get('lastChangeAll'), sessions=sessions),
          open(os.path.join(a.out, 'snapshot.json'), 'w', encoding='utf-8'), ensure_ascii=False)

# --- ICS ---
def esc(s): return s.replace('\\', '\\\\').replace(';', '\\;').replace(',', '\\,').replace('\n', '\\n')
def fold(line):
    b = line.encode('utf-8'); out = []
    while len(b) > 73:
        cut = 73
        while (b[cut] & 0xC0) == 0x80: cut -= 1
        out.append(b[:cut].decode()); b = b' ' + b[cut:]
    out.append(b.decode()); return '\r\n'.join(out)
I = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//ilhan//lodz-takip//TR", "CALSCALE:GREGORIAN", "X-WR-CALNAME:Łódź Dersler", "X-WR-TIMEZONE:Europe/Warsaw",
     "BEGIN:VTIMEZONE", "TZID:Europe/Warsaw", "BEGIN:STANDARD", "DTSTART:19701025T030000", "TZOFFSETFROM:+0200", "TZOFFSETTO:+0100", "RRULE:FREQ=YEARLY;BYMONTH=10;BYDAY=-1SU", "END:STANDARD",
     "BEGIN:DAYLIGHT", "DTSTART:19700329T020000", "TZOFFSETFROM:+0100", "TZOFFSETTO:+0200", "RRULE:FREQ=YEARLY;BYMONTH=3;BYDAY=-1SU", "END:DAYLIGHT", "END:VTIMEZONE"]
stamp = dt.datetime.utcnow().strftime('%Y%m%dT%H%M%SZ')
for x in sessions:
    d = dt.date.fromisoformat(x['date'])
    st = dt.datetime.combine(d, dt.time()) + dt.timedelta(minutes=x['s']); en = dt.datetime.combine(d, dt.time()) + dt.timedelta(minutes=x['e'])
    grp = 'herkes' if x['g'] == 'herkes' else 'Grup: ' + x['g']
    extra = f" – {x['n']}" if x['n'] and x['t'] != 'Lecture' else ''
    desc = f"Grubun: {x['g']}\nCelcat'ta yazan: {x['cel']}\nHoca: {x['p']}" + (f"\nNot: {x['n']}" if x['n'] else '')
    for ln in ["BEGIN:VEVENT", f"UID:lodz-{x['id']}-{x['date']}@ilhan", f"DTSTAMP:{stamp}", f"DTSTART;TZID=Europe/Warsaw:{st:%Y%m%dT%H%M%S}", f"DTEND;TZID=Europe/Warsaw:{en:%Y%m%dT%H%M%S}",
               "SUMMARY:" + esc(f"{x['c']} {x['t']} – {grp}{extra}"), "LOCATION:" + esc(x['r']), "DESCRIPTION:" + esc(desc), "END:VEVENT"]:
        I.append(fold(ln))
I.append("END:VCALENDAR")
open(os.path.join(a.out, 'lodz-ders-programi.ics'), 'w', encoding='utf-8', newline='').write('\r\n'.join(I) + '\r\n')

# --- Excel ---
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
COL = {'CC': 'DCE8FB', 'IoT': 'D8F1E8', 'ML': 'E9E1FA', 'WC': 'FDEBD0', 'DS': 'E1E8EF', 'PMM': 'EEF1D5', 'IT': 'FFE3D3'}
F = Font(name='Arial', size=10); FB = Font(name='Arial', size=10, bold=True, color='FFFFFF')
HF = PatternFill('solid', fgColor='8A1C1C'); B = Border(bottom=Side(style='thin', color='D0D5DD'))
wb = Workbook()
def header(ws, cols, widths):
    ws.append(cols)
    for i, c in enumerate(ws[1], 1):
        c.font = FB; c.fill = HF; ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = 'A2'
ws = wb.active; ws.title = 'Değişiklikler'
header(ws, ['Kontrol', 'Durum'], [22, 120])
ws.append([dt.datetime.now().strftime('%d.%m.%Y %H:%M'), L[0]])
for typ, _, _, txt in changes:
    ws.append(['', f"{typ}: {txt}"])
for x, y in conflicts:
    ws.append(['', f"ÇAKIŞMA: {label(x)} ⟷ {y['c']} {y['t']} {hm(y['s'])}–{hm(y['e'])}"])
for row in ws.iter_rows(min_row=2):
    for c in row: c.font = F; c.alignment = Alignment(wrap_text=True, vertical='top')
ws2 = wb.create_sheet('Haftalık Şablon')
header(ws2, ['Ders', 'Tür', 'Grup', 'Gün', 'Başlangıç', 'Bitiş', 'Salon', 'Hoca', 'Not', 'Tarihler'], [30, 12, 34, 11, 10, 8, 26, 26, 26, 60])
byev = defaultdict(list)
for x in sessions: byev[x['id']].append(x)
for eid, xs in sorted(byev.items(), key=lambda kv: (dt.date.fromisoformat(kv[1][0]['date']).weekday(), kv[1][0]['s'])):
    x = xs[0]
    ws2.append([x['c'], x['t'], x['g'], GUN[dt.date.fromisoformat(x['date']).weekday()], hm(x['s']), hm(x['e']), x['r'], x['p'], x['n'],
                ', '.join(dt.date.fromisoformat(y['date']).strftime('%d.%m') for y in xs)])
    for c in ws2[ws2.max_row]: c.font = F; c.fill = PatternFill('solid', fgColor=COL[x['k']]); c.border = B; c.alignment = Alignment(vertical='top', wrap_text=(c.column == 10))
ws2.auto_filter.ref = ws2.dimensions
ws3 = wb.create_sheet('Tüm Dersler')
header(ws3, ['Tarih', 'Gün', 'Başlangıç', 'Bitiş', 'Ders', 'Tür', 'Grup', 'Salon', 'Hoca', 'Not'], [12, 11, 10, 8, 30, 12, 34, 26, 26, 26])
for x in sessions:
    d = dt.date.fromisoformat(x['date'])
    ws3.append([d, GUN[d.weekday()], hm(x['s']), hm(x['e']), x['c'], x['t'], x['g'], x['r'], x['p'], x['n']])
    r = ws3[ws3.max_row]
    for c in r: c.font = F; c.fill = PatternFill('solid', fgColor=COL[x['k']]); c.border = B
    r[0].number_format = 'DD.MM.YYYY'
ws3.auto_filter.ref = ws3.dimensions
ws4 = wb.create_sheet('Grup Seçimleri')
header(ws4, ['Ders', 'Tür', 'Celcat grubu', 'Gün / saat', 'Salon', 'Hangi haftalar', 'Seçildi mi', 'Neden'], [28, 11, 34, 26, 14, 30, 11, 70])
for rr in cfg.get('grup_secimleri', []):
    ws4.append(rr)
    for c in ws4[ws4.max_row]:
        c.font = Font(name='Arial', size=10, bold=(rr[6] == '✅')); c.border = B; c.alignment = Alignment(vertical='top', wrap_text=True)
ws5 = wb.create_sheet('Notlar')
for n in cfg.get('notes', []) + [['Son kontrol', dt.datetime.now().strftime('%d.%m.%Y %H:%M')], ['Celcat son değişiklik', str(raw.get('lastChangeAll', ''))[:16]]] + \
         [['Programdan çıkarılan', f"{k[1]} ({k[0]}): {v}"] for k, v in skip.items()]:
    ws5.append(n)
ws5.column_dimensions['A'].width = 22; ws5.column_dimensions['B'].width = 110
for row in ws5.iter_rows(): row[0].font = Font(name='Arial', size=10, bold=True); row[1].font = F
wb.save(os.path.join(a.out, 'lodz-ders-programi.xlsx'))

# --- HTML sayfası ---
tpl = open(a.template, encoding='utf-8').read()
D = [dict(c=C[e['k']], k=e['k'], t=e['t'], g=e['g'] + (f" · {e['n'].strip()}" if e['n'].strip() and e['t'] != 'Lecture' else ''), d=e['d'], s=e['s'], u=e['u'], w=e['w'],
          r=room(e['r'], e['n']), p=e['p']) for e in events if '1' in e['w']]
if conflicts:
    st = '<div><b>Çakışma var:</b> ' + '; '.join(f"{tr_date(dt.date.fromisoformat(x['date']))} {x['c']} ⟷ {y['c']}" for x, y in conflicts) + '</div>'
else:
    st = '<div style="border-left-color:var(--c-IoT)"><b>7 ders, 24 ECTS. Bugünden sonra çakışma yok.</b></div>'
if changes:
    st += '<div style="border-left-color:var(--accent)"><b>Son değişiklikler:</b><br>' + '<br>'.join(f"{t}: {txt}" for t, _, _, txt in changes) + '</div>'
for k, v in skip.items():
    st += f'<div style="border-left-color:var(--muted)">{k[1]}: {C.get(next((e["k"] for e in raw["events"] if e["id"] == k[0]), ""), "")} programdan çıkarıldı – {v}</div>'
page = tpl.replace('__DATA__', json.dumps(D, ensure_ascii=False)).replace('__STATUS__', st).replace('__CHECKED__', dt.datetime.now().strftime('%d.%m.%Y'))
open(os.path.join(a.out, 'page.html'), 'w', encoding='utf-8').write(page)

res = dict(changed=bool(changes), n_changes=len(changes), conflicts=len(conflicts), sessions=len(sessions), first_run=prev is None)
json.dump(res, open(os.path.join(a.out, 'sonuc.json'), 'w'), ensure_ascii=False)
print(json.dumps(res, ensure_ascii=False))
print(report)
